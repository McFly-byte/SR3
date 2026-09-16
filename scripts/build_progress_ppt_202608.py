from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"D:\LMC\projects\Image-Super-Resolution-via-Iterative-Refinement")
OUT_DIR = ROOT / "paper_material" / "20260814组会" / "课题当前进展"
OUT_FILE = OUT_DIR / "20260817组会.pptx"

WATER_ROOT = Path(
    r"D:\LMC\水膜数据处理\20260716后-水膜\phantom_quantification_revision"
    r"\定量验证\第六阶段_最终结果包\run_20260725_165651"
)
SIM_ROOT = Path(
    r"D:\LMC\data\simulated_with_lesion\0716后-仿真\第十阶段_case_QC"
    r"\sub-03\lesion003\mixed_level2_s0375_lr24"
)
EVIDENCE_ROOT = OUT_DIR / "推理对比材料" / "汇总"

IMG_WATER_FIX = WATER_ROOT / "09_previous_problem_reproduction_and_fix.png"
IMG_WATER_COMPARE = WATER_ROOT / "06_senior_low_high_relative_quantification_comparison.png"
IMG_SIM_QC = SIM_ROOT / "第十阶段_case_QC总览.png"
IMG_SIM_CURVES = SIM_ROOT / "第十阶段_四区域动态曲线_QC.png"
IMG_SIM_COMPARE = (
    ROOT
    / "experiments"
    / "analyses"
    / "training_diagnostics_20260730"
    / "improvement_runs"
    / "qualitative"
    / "03_baseline_vs_selected_viridis.png"
)
IMG_INVIVO_117 = EVIDENCE_ROOT / "活体_mouse_11.7T_四代谢物_微调前后.png"
IMG_INVIVO_94 = EVIDENCE_ROOT / "活体_mouse_9.4T_四代谢物_微调前后.png"


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(247, 248, 250)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 39, 49)
MUTED = RGBColor(85, 98, 110)
LINE = RGBColor(216, 222, 228)
TEAL = RGBColor(0, 124, 131)
BLUE = RGBColor(47, 107, 255)
GREEN = RGBColor(42, 157, 111)
RED = RGBColor(216, 74, 58)
AMBER = RGBColor(225, 150, 42)
PALE_TEAL = RGBColor(229, 244, 244)
PALE_BLUE = RGBColor(234, 239, 255)
PALE_GREEN = RGBColor(232, 246, 239)
PALE_RED = RGBColor(252, 235, 232)
PALE_AMBER = RGBColor(251, 242, 226)

FONT_CN = "Microsoft YaHei"
FONT_EN = "Aptos"


def set_background(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    font=FONT_CN,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_rich_text(slide, x, y, w, h, runs, size=17, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    for spec in runs:
        r = p.add_run()
        r.text = spec[0]
        r.font.name = FONT_CN
        r.font.size = Pt(spec[1] if len(spec) > 1 else size)
        r.font.bold = spec[2] if len(spec) > 2 else False
        r.font.color.rgb = spec[3] if len(spec) > 3 else INK
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, 0.55, 0.28, 12.25, 0.48, title, size=26, color=INK, bold=True)
    if subtitle:
        add_text(slide, 0.58, 0.76, 12.0, 0.30, subtitle, size=11.5, color=MUTED)


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(0.8)
    return shp


def add_metric(slide, x, y, w, h, value, label, accent=TEAL, note=None, value_size=24):
    add_rect(slide, x, y, w, h, WHITE, LINE)
    add_text(slide, x + 0.13, y + 0.12, w - 0.26, 0.43, value, size=value_size, color=accent, bold=True)
    add_text(slide, x + 0.13, y + 0.57, w - 0.26, 0.30, label, size=11.5, color=INK, bold=True)
    if note:
        add_text(slide, x + 0.13, y + 0.88, w - 0.26, h - 0.96, note, size=9.5, color=MUTED)


def add_bullet_list(slide, x, y, w, h, items, size=14, color=INK, bullet_color=TEAL, gap=0.12):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(gap * 10)
        p.line_spacing = 1.12
        r0 = p.add_run()
        r0.text = "●  "
        r0.font.name = FONT_CN
        r0.font.size = Pt(max(8, size - 3))
        r0.font.color.rgb = bullet_color
        r1 = p.add_run()
        r1.text = item
        r1.font.name = FONT_CN
        r1.font.size = Pt(size)
        r1.font.color.rgb = color
    return box


def add_picture_contain(slide, path, x, y, w, h, line=True):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    if line:
        add_rect(slide, x, y, w, h, WHITE, LINE, radius=False)
    return slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def add_picture_cover(slide, path, x, y, w, h, line=True):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    target_ratio = w / h
    img_ratio = iw / ih
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if img_ratio > target_ratio:
        visible = target_ratio / img_ratio
        crop = (1 - visible) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        visible = img_ratio / target_ratio
        crop = (1 - visible) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    if line:
        frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        frame.fill.background()
        frame.line.color.rgb = LINE
        frame.line.width = Pt(0.8)
    return pic


def add_step(slide, x, y, w, h, number, title, detail, color, pale):
    add_rect(slide, x, y, w, h, WHITE, LINE)
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(y + 0.16), Inches(0.42), Inches(0.42))
    circ.fill.solid()
    circ.fill.fore_color.rgb = color
    circ.line.color.rgb = color
    add_text(slide, x + 0.15, y + 0.16, 0.42, 0.42, str(number), size=13, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, x + 0.68, y + 0.12, w - 0.82, 0.35, title, size=15, color=INK, bold=True)
    add_text(slide, x + 0.16, y + 0.67, w - 0.32, h - 0.76, detail, size=10.7, color=MUTED)


def add_arrow(slide, x, y, w=0.32, h=0.16, color=LINE):
    shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.color.rgb = color
    return shp


def add_table(slide, x, y, w, h, data, col_widths=None, header_fill=INK, font_size=11.5):
    rows, cols = len(data), len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else RGBColor(242, 245, 247))
            cell.border = None
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = FONT_CN
                    run.font.size = Pt(font_size)
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
    return table


def build_deck():
    for path in [IMG_WATER_FIX, IMG_WATER_COMPARE, IMG_SIM_QC, IMG_SIM_CURVES,
                 IMG_SIM_COMPARE, IMG_INVIVO_117, IMG_INVIVO_94]:
        if not path.exists():
            raise FileNotFoundError(path)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # Slide 1: overview
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_text(slide, 0.66, 0.56, 11.9, 0.82, "过去一个半月工作进展", size=34, color=INK, bold=True)
    add_text(slide, 0.70, 1.43, 11.8, 0.42,
             "扩散生成先验与结构引导的氘代谢 MRSI 超分辨率重建", size=18, color=TEAL, bold=True)
    add_text(slide, 0.70, 1.93, 11.5, 0.32, "数据处理纠错 · 病灶仿真 · 模型微调 · 活体验证", size=13, color=MUTED)

    add_step(slide, 0.68, 2.68, 2.78, 2.62, 1, "水膜处理纠错",
             "修正数据域、群延迟、方向\n和无依据浓度标定", TEAL, PALE_TEAL)
    add_arrow(slide, 3.53, 3.85)
    add_step(slide, 3.90, 2.68, 2.78, 2.62, 2, "病灶仿真闭环",
             "建立动态代谢、病灶配对、\nk-space 退化与质量控制", BLUE, PALE_BLUE)
    add_arrow(slide, 6.75, 3.85)
    add_step(slide, 7.12, 2.68, 2.78, 2.62, 3, "训练与保真微调",
             "完成 50k 主训练\n和 500-step 短程微调", GREEN, PALE_GREEN)
    add_arrow(slide, 9.97, 3.85)
    add_step(slide, 10.34, 2.68, 2.32, 2.62, 4, "活体验证",
             "真实小鼠暴露跨物种、\n跨矩阵和结构条件偏移", RED, PALE_RED)

    add_rect(slide, 0.68, 5.72, 11.98, 0.82, PALE_TEAL, PALE_TEAL)
    add_rich_text(slide, 0.92, 5.89, 11.50, 0.45, [
        ("核心进展：", 16, True, TEAL),
        ("从“图像更清晰”推进到数据域、定量和采集一致性可验证", 16, True, INK),
    ], valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.70, 6.86, 6.0, 0.28, "组会汇报 · 2026-08", size=11, color=MUTED)

    # Slide 2: water phantom
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "1  水膜处理：早期问题与最新修正", "真实 MRSI 处理链先于超分模型验证")
    add_picture_contain(slide, IMG_WATER_FIX, 0.55, 1.18, 7.55, 5.82)

    add_rect(slide, 8.35, 1.18, 4.43, 2.43, WHITE, LINE)
    add_text(slide, 8.57, 1.37, 3.98, 0.36, "早期版本的关键问题", size=16, color=RED, bold=True)
    add_bullet_list(slide, 8.55, 1.82, 4.0, 1.56, [
        "对逐体素图像域数据重复做空间 IFFT",
        "fid_proc.64 分支遗漏约 77 点群延迟",
        "转置与 Y 轴约定不一致",
        "无外部标定却输出 mM，自动色标放大异常",
    ], size=11.7, bullet_color=RED, gap=0.05)

    add_rect(slide, 8.35, 3.82, 4.43, 1.72, WHITE, LINE)
    add_text(slide, 8.57, 4.00, 3.98, 0.34, "新版处理原则", size=16, color=TEAL, bold=True)
    add_bullet_list(slide, 8.55, 4.40, 4.0, 0.94, [
        "只沿时间维 FFT，不再额外空间 IFFT",
        "物理坐标统一结构像与谱图",
        "报告相对峰面积，不冒充绝对浓度",
    ], size=11.7, bullet_color=TEAL, gap=0.04)

    add_metric(slide, 8.35, 5.78, 2.09, 1.21, "104.63%", "本人全图峰面积和比", GREEN)
    add_metric(slide, 10.69, 5.78, 2.09, 1.21, "102.30%", "师兄全图峰面积和比", GREEN)

    # Slide 3: simulation
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "2  病灶仿真：从结构体模到可重放低分辨率观测")

    flow_titles = ["HBA", "组织概率", "动态代谢", "病灶建模", "HR 配对", "k-space", "训练样本"]
    flow_colors = [TEAL, TEAL, BLUE, RED, GREEN, AMBER, BLUE]
    x0, fw, gap = 0.56, 1.54, 0.26
    for i, (name, color) in enumerate(zip(flow_titles, flow_colors)):
        x = x0 + i * (fw + gap)
        add_rect(slide, x, 1.05, fw, 0.64, WHITE, color)
        add_text(slide, x + 0.02, 1.10, fw - 0.04, 0.53, name, size=12.5, color=color, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(flow_titles) - 1:
            add_arrow(slide, x + fw + 0.03, 1.30, w=0.20, h=0.13, color=LINE)

    add_picture_contain(slide, IMG_SIM_QC, 0.55, 1.93, 8.32, 5.10)

    add_metric(slide, 9.10, 1.93, 1.67, 1.18, "17", "动态时间点", BLUE)
    add_metric(slide, 10.91, 1.93, 1.87, 1.18, "4", "HDO/Glc/Glx/Lac", TEAL)
    add_metric(slide, 9.10, 3.32, 1.67, 1.18, "3", "LR 16/24/32", AMBER)
    add_metric(slide, 10.91, 3.32, 1.87, 1.18, "204", "可逐元素重放", GREEN)

    add_rect(slide, 9.10, 4.75, 3.68, 2.28, WHITE, LINE)
    add_text(slide, 9.32, 4.95, 3.25, 0.34, "改进版 v1 验收", size=16, color=INK, bold=True)
    add_bullet_list(slide, 9.30, 5.40, 3.15, 1.34, [
        "64×64 RAS 固定网格，FOV 216 mm",
        "健康/异常严格反事实配对",
        "复高斯 k-space 噪声与完整 metadata",
        "23/23 单元测试通过",
    ], size=11.5, bullet_color=GREEN, gap=0.04)

    # Slide 4: paired simulation comparison
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(
        slide,
        "3  微调改善仿真平均误差，但病灶场景仍落后健康约 6 dB",
        "同一批 408 个 validation 样本：healthy 204、lesion 204；I50000 与 I50500 使用 raw 网络",
    )
    add_picture_contain(slide, IMG_SIM_COMPARE, 0.55, 1.16, 6.91, 4.90)

    healthy_table = [
        ["健康 n=204", "I50000", "I50500", "相对变化"],
        ["PSNR / dB", "39.15", "39.63", "+1.2%"],
        ["masked PSNR", "36.93", "37.73", "+2.2%"],
        ["masked MAE", "0.0141", "0.0118", "−16.3%"],
        ["ROI mean error", "4.48%", "2.96%", "−33.9%"],
    ]
    lesion_table = [
        ["病灶 n=204", "I50000", "I50500", "相对变化"],
        ["PSNR / dB", "33.27", "33.57", "+0.9%"],
        ["masked PSNR", "29.87", "30.31", "+1.5%"],
        ["masked MAE", "0.0215", "0.0194", "−9.6%"],
        ["ROI mean error", "5.86%", "5.12%", "−12.7%"],
    ]
    add_table(slide, 7.70, 1.16, 5.08, 2.20, healthy_table,
              col_widths=[1.70, 1.03, 1.03, 1.32], header_fill=TEAL, font_size=9.7)
    add_table(slide, 7.70, 3.61, 5.08, 2.20, lesion_table,
              col_widths=[1.70, 1.03, 1.03, 1.32], header_fill=RED, font_size=9.7)

    add_rect(slide, 0.55, 6.25, 12.23, 0.78, PALE_RED, PALE_RED)
    add_rich_text(slide, 0.78, 6.38, 11.77, 0.47, [
        ("微调后差距：", 13.5, True, RED),
        ("健康与病灶 PSNR 相差 6.06 dB；总体 SSIM 反而下降 0.0044。", 13.5, True, INK),
        (" 病灶域仍不是可靠优势场景。", 13.5, True, RED),
    ], valign=MSO_ANCHOR.MIDDLE)

    # Slide 5: 11.7 T in-vivo paired comparison
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "4  11.7 T 活体小鼠：微调前后近似，无 HR 时不能判定病灶真实性")
    add_picture_contain(slide, IMG_INVIVO_117, 0.55, 1.05, 9.48, 5.98)

    add_rect(slide, 10.25, 1.05, 2.53, 1.40, WHITE, LINE)
    add_text(slide, 10.47, 1.25, 2.09, 0.30, "比较条件", size=15, color=INK, bold=True)
    add_text(slide, 10.47, 1.66, 2.05, 0.60,
             "9×9 · 5 层 · 4 代谢物\n同一 LR / scale / 5 seeds",
             size=10.6, color=MUTED)

    add_rect(slide, 10.25, 2.68, 2.53, 2.36, WHITE, LINE)
    add_text(slide, 10.47, 2.87, 2.09, 0.30, "全 20 样本观测指标", size=14, color=INK, bold=True)
    add_text(slide, 10.47, 3.35, 2.07, 1.47,
             "回算相对 L1\n0.636 → 0.623  (−2.1%)\n\n背景泄漏\n0.287 → 0.281  (−2.1%)\n\n梯度比\n1.637 → 1.632  (−0.3%)",
             size=10.3, color=INK)

    add_rect(slide, 10.25, 5.28, 2.53, 1.75, PALE_RED, PALE_RED)
    add_text(slide, 10.47, 5.47, 2.08, 1.32,
             "两版输出和观测指标均接近；无 HR 真值，不能把新增纹理解释为真实病灶。",
             size=10.8, color=RED, bold=True, valign=MSO_ANCHOR.MIDDLE)

    # Slide 6: 9.4 T in-vivo paired comparison
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "5  9.4 T 活体小鼠：7×7 分布外输入下，微调未改变主要空间模式")
    add_picture_contain(slide, IMG_INVIVO_94, 0.55, 1.05, 9.48, 5.98)

    add_rect(slide, 10.25, 1.05, 2.53, 1.40, WHITE, LINE)
    add_text(slide, 10.47, 1.25, 2.09, 0.30, "比较条件", size=15, color=INK, bold=True)
    add_text(slide, 10.47, 1.66, 2.05, 0.60,
             "7×7 · 3 层 · 4 代谢物\n同一 LR / scale / 5 seeds",
             size=10.6, color=MUTED)

    add_rect(slide, 10.25, 2.68, 2.53, 2.36, WHITE, LINE)
    add_text(slide, 10.47, 2.87, 2.09, 0.30, "全 12 样本观测指标", size=14, color=INK, bold=True)
    add_text(slide, 10.47, 3.35, 2.07, 1.47,
             "回算相对 L1\n0.405 → 0.405  (−0.1%)\n\n背景泄漏\n0.380 → 0.375  (−1.3%)\n\n梯度比\n1.792 → 1.793  (+0.1%)",
             size=10.3, color=INK)

    add_rect(slide, 10.25, 5.28, 2.53, 1.75, PALE_RED, PALE_RED)
    add_text(slide, 10.47, 5.47, 2.08, 1.32,
             "7×7 低于训练的 16/24/32 矩阵。现阶段只能报告输出形态，不能声称病灶恢复正确。",
             size=11.4, color=RED, bold=True, valign=MSO_ANCHOR.MIDDLE)

    # Slide 7: conclusion and next plan
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "6  当前证据不支持直接解释病灶；下一步先完成健康小鼠域适配")

    add_rect(slide, 0.55, 1.10, 4.00, 2.08, WHITE, LINE)
    add_text(slide, 0.80, 1.34, 3.52, 0.35, "为什么先做健康小鼠", size=18, color=TEAL, bold=True)
    add_bullet_list(slide, 0.77, 1.86, 3.52, 1.03, [
        "暂时移除病灶结构/代谢匹配这一额外变量",
        "先验证真实小鼠解剖、矩阵和结构条件",
        "健康脑不稳定时，病灶结论没有可靠基础",
    ], size=12.2, bullet_color=TEAL, gap=0.07)

    add_rect(slide, 4.80, 1.10, 3.76, 2.08, PALE_BLUE, PALE_BLUE)
    add_text(slide, 5.06, 1.34, 3.25, 0.35, "有配对高/低分辨率", size=17, color=BLUE, bold=True)
    add_bullet_list(slide, 5.04, 1.86, 3.22, 1.04, [
        "核对 FOV、层厚、时间点和相对量纲",
        "开展监督或半监督小学习率微调",
        "用原始矩阵回算约束选模",
    ], size=11.8, bullet_color=BLUE, gap=0.06)

    add_rect(slide, 8.81, 1.10, 3.97, 2.08, PALE_AMBER, PALE_AMBER)
    add_text(slide, 9.07, 1.34, 3.45, 0.35, "只有低分辨率数据", size=17, color=AMBER, bold=True)
    add_bullet_list(slide, 9.05, 1.86, 3.42, 1.04, [
        "不能把插值结果当作 HR 标签",
        "先做低矩阵模拟适配和采集一致性正则",
        "增加背景约束与结构反事实测试",
    ], size=11.8, bullet_color=AMBER, gap=0.06)

    stages = [
        ("01", "数据审计", "设备、矩阵、数据域、结构模态\n采集参数、哈希与频谱 QC", TEAL),
        ("02", "健康脑基线", "bicubic / 当前模型\n回算、背景、多 seed、热点外推", BLUE),
        ("03", "真实域适配", "覆盖真实 7×7 / 9×9\n或师姐数据的实际矩阵", AMBER),
        ("04", "通过后回到病灶", "结构病灶一致仿真\nlesion-aware loss 与病灶评价", GREEN),
    ]
    y = 3.68
    for i, (num, title, detail, color) in enumerate(stages):
        x = 0.55 + i * 3.08
        add_rect(slide, x, y, 2.78, 2.27, WHITE, LINE)
        add_text(slide, x + 0.17, y + 0.15, 0.55, 0.35, num, size=15, color=color, bold=True)
        add_text(slide, x + 0.17, y + 0.55, 2.42, 0.38, title, size=17, color=INK, bold=True)
        add_text(slide, x + 0.17, y + 1.06, 2.42, 0.91, detail, size=10.7, color=MUTED)
        if i < 3:
            add_arrow(slide, x + 2.84, y + 1.07, w=0.18, h=0.14, color=LINE)

    add_rect(slide, 0.55, 6.30, 12.23, 0.75, PALE_GREEN, PALE_GREEN)
    add_text(slide, 0.82, 6.42, 11.70, 0.46,
             "阶段目标：在多例健康小鼠上同时满足采集一致性、背景控制和结构反事实稳定，再进入病灶模型。",
             size=14.5, color=GREEN, bold=True, valign=MSO_ANCHOR.MIDDLE)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build_deck()
